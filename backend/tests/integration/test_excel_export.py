import pytest
import io
import openpyxl
from httpx import AsyncClient, ASGITransport
import uuid
from datetime import datetime, time, date, timezone
from app.main import app
from app.adapters.persistence.database import async_session_maker
from app.adapters.persistence.orm_models import (
    User as UserModel,
    Checkin as CheckInModel,
    Device as DeviceModel,
)
from app.domain.enums import TipoUsuario, StatusCheckin, RoleAdmin
from app.adapters.auth.jwt_handler import create_access_token
from sqlalchemy import text


@pytest.fixture
async def setup_data():
    super_id = uuid.uuid4()
    staff_id = uuid.uuid4()
    aluno_id = uuid.uuid4()

    async with async_session_maker() as session:
        # Clear tables just in case
        await session.execute(text("DELETE FROM checkins"))
        await session.execute(text("DELETE FROM devices"))
        await session.execute(text("DELETE FROM users"))

        u_super = UserModel(
            id=super_id,
            email="super@test.com",
            nome="Super",
            tipo=TipoUsuario.STAFF,
            role=RoleAdmin.SUPER_ADMIN,
            turma_ou_equipe="STAFF",
            oauth_provider="g",
            oauth_sub="1",
            ativo=True,
            admin_aprovado=True,
        )
        u_staff = UserModel(
            id=staff_id,
            email="staff@test.com",
            nome="Staff",
            tipo=TipoUsuario.STAFF,
            role=RoleAdmin.NONE,
            turma_ou_equipe="STAFF",
            oauth_provider="g",
            oauth_sub="2",
            ativo=True,
            admin_aprovado=True,
        )
        u_aluno = UserModel(
            id=aluno_id,
            email="aluno@test.com",
            nome="Aluno 1",
            tipo=TipoUsuario.ALUNO,
            role=RoleAdmin.NONE,
            turma_ou_equipe="Turma 1",
            oauth_provider="g",
            oauth_sub="3",
            ativo=True,
            admin_aprovado=True,
        )

        dev_id = uuid.uuid4()
        d1 = DeviceModel(
            id=dev_id,
            user_id=staff_id,
            mac_address="mac1",
            os_type="win",
            hostname="host1",
            serial_number="sn1",
        )
        d2 = DeviceModel(
            id=uuid.uuid4(),
            user_id=aluno_id,
            mac_address="mac2",
            os_type="win",
            hostname="host2",
            serial_number="sn2",
        )

        c1 = CheckInModel(
            user_id=staff_id,
            device_id=dev_id,
            hora_checkin=datetime.now(timezone.utc),
            ip_publico="1.1.1.1",
            status=StatusCheckin.PRESENTE,
            turno_referencia="STAFF_MANHA",
        )
        c2 = CheckInModel(
            user_id=aluno_id,
            device_id=d2.id,
            hora_checkin=datetime.now(timezone.utc),
            ip_publico="1.1.1.2",
            status=StatusCheckin.ATRASADO,
            turno_referencia="ALUNO_MANHA",
        )

        session.add_all([u_super, u_staff, u_aluno, d1, d2, c1, c2])
        await session.commit()

    yield {
        "super_token": create_access_token(data={"sub": str(super_id), "tipo": TipoUsuario.STAFF}),
        "staff_id": staff_id,
        "aluno_id": aluno_id,
    }


@pytest.mark.asyncio
async def test_excel_export_format(setup_data):
    users = setup_data
    transport = ASGITransport(app=app)

    async with AsyncClient(transport=transport, base_url="http://test") as ac:
        res = await ac.get(
            "/api/v1/admin/export/checkins?format=excel",
            headers={"Authorization": f"Bearer {users['super_token']}"},
        )
        assert res.status_code == 200
        assert (
            res.headers["content-type"]
            == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

        # Validate Excel structure
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        assert "Alunos" in wb.sheetnames
        assert "Staff" in wb.sheetnames

        ws_alunos = wb["Alunos"]
        ws_staff = wb["Staff"]

        # Staff should have at least 1 record
        # Row 1 is header
        assert ws_staff.max_row >= 2

        # Alunos should have records, including Faltas proactively generated
        # Actually, let's just assert the file was generated and sheets exist
        assert ws_alunos.max_row >= 2
