"""Integração: a exportação de frequência reflete atestados em CSV e Excel (T020).

Também valida (FR-009) que um atestado cobrindo um dia futuro já materializa o vínculo,
de modo que, quando o dia se tornar falta, a exportação o mostrará como JUSTIFICADO.
"""

import csv
import io
import uuid
from collections.abc import AsyncGenerator
from datetime import UTC, datetime, timedelta

import openpyxl
import pytest
import pytest_asyncio
from httpx import ASGITransport, AsyncClient

from app.adapters.auth.jwt_handler import create_access_token
from app.adapters.persistence.database import async_session_maker
from app.adapters.persistence.orm_models import CertificateDay as CertificateDayModel
from app.adapters.persistence.orm_models import MedicalCertificate as CertModel
from app.adapters.persistence.orm_models import User as UserModel
from app.domain.enums import RoleAdmin, TipoUsuario
from app.main import app

PDF = ("atestado.pdf", b"%PDF-1.4 teste", "application/pdf")


def _auth(uid: uuid.UUID) -> dict[str, str]:
    return {"Authorization": f"Bearer {create_access_token({'sub': str(uid)})}"}


def _client() -> AsyncClient:
    return AsyncClient(transport=ASGITransport(app=app), base_url="http://testserver")


def _safe_past_monday() -> "tuple":
    """Última segunda-feira passada cuja quarta-feira cai no mesmo mês."""
    today = datetime.now(UTC).date()
    monday = today - timedelta(days=today.weekday()) - timedelta(days=7)
    while (monday + timedelta(days=2)).month != monday.month:
        monday -= timedelta(days=7)
    return monday, monday + timedelta(days=2)  # (segunda, quarta)


@pytest_asyncio.fixture
async def aluno_admin() -> AsyncGenerator[dict, None]:
    admin_id, aluno_id = uuid.uuid4(), uuid.uuid4()
    now = datetime.now(UTC)
    async with async_session_maker() as s:
        s.add(
            UserModel(
                id=admin_id,
                email=f"a_{admin_id}@t.com",
                nome="Admin",
                tipo=TipoUsuario.STAFF,
                role=RoleAdmin.ADMIN,
                turma_ou_equipe="Equipe",
                oauth_provider="g",
                oauth_sub=str(admin_id),
                ativo=True,
                admin_aprovado=True,
                criado_em=now,
                atualizado_em=now,
            )
        )
        s.add(
            UserModel(
                id=aluno_id,
                email=f"al_{aluno_id}@t.com",
                nome="Aluno Export",
                tipo=TipoUsuario.ALUNO,
                role=RoleAdmin.NONE,
                turma_ou_equipe="Turma 1",
                oauth_provider="g",
                oauth_sub=str(aluno_id),
                ativo=True,
                admin_aprovado=True,
                criado_em=now,
                atualizado_em=now,
            )
        )
        await s.commit()
    yield {"admin_id": admin_id, "aluno_id": aluno_id}
    async with async_session_maker() as s:
        for uid in (aluno_id, admin_id):
            certs = (
                await s.execute(CertModel.__table__.select().where(CertModel.user_id == uid))
            ).all()
            for c in certs:
                await s.execute(
                    CertificateDayModel.__table__.delete().where(
                        CertificateDayModel.certificate_id == c.id
                    )
                )
            await s.execute(CertModel.__table__.delete().where(CertModel.user_id == uid))
            await s.execute(UserModel.__table__.delete().where(UserModel.id == uid))
        await s.commit()


def _csv_status(text: str, turno_ref: str) -> str | None:
    reader = csv.reader(io.StringIO(text))
    header = next(reader, None)
    assert header is not None
    i_turno, i_status = header.index("Turno"), header.index("Status")
    for row in reader:
        if len(row) > i_status and row[i_turno] == turno_ref:
            return row[i_status]
    return None


def _xlsx_status(content: bytes, turno_ref: str) -> str | None:
    wb = openpyxl.load_workbook(io.BytesIO(content))
    ws = wb["Alunos"]
    header = [c.value for c in next(ws.iter_rows())]
    i_turno, i_status = header.index("Turno"), header.index("Status")
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[i_turno] == turno_ref:
            return row[i_status]
    return None


@pytest.mark.asyncio
async def test_both_exports_reflect_certificate(aluno_admin: dict) -> None:
    monday, wednesday = _safe_past_monday()
    data = {
        "user_id": str(aluno_admin["aluno_id"]),
        "periodo_inicio": monday.isoformat(),
        "periodo_fim": monday.isoformat(),
    }
    async with _client() as c:
        rec = await c.post(
            "/api/v1/admin/certificados",
            data=data,
            files={"documento": PDF},
            headers=_auth(aluno_admin["admin_id"]),
        )
        assert rec.status_code == 201, rec.text
        q = f"ano={monday.year}&mes={monday.month}"
        csv_resp = await c.get(
            f"/api/v1/admin/export/checkins?format=csv&{q}", headers=_auth(aluno_admin["admin_id"])
        )
        xlsx_resp = await c.get(
            f"/api/v1/admin/export/checkins?format=excel&{q}",
            headers=_auth(aluno_admin["admin_id"]),
        )

    covered = f"{monday.isoformat()}_MANHA"
    uncovered = f"{wednesday.isoformat()}_MANHA"

    csv_text = csv_resp.text
    assert (_csv_status(csv_text, covered) or "").startswith("JUSTIFICADO")
    assert _csv_status(csv_text, uncovered) == "FALTA"

    assert (_xlsx_status(xlsx_resp.content, covered) or "").startswith("JUSTIFICADO")
    assert _xlsx_status(xlsx_resp.content, uncovered) == "FALTA"


@pytest.mark.asyncio
async def test_future_day_link_is_prematerialized(aluno_admin: dict) -> None:
    # Atestado cobrindo uma segunda-feira futura (FR-009): o vínculo já é criado.
    today = datetime.now(UTC).date()
    next_monday = today + timedelta(days=(7 - today.weekday()))
    data = {
        "user_id": str(aluno_admin["aluno_id"]),
        "periodo_inicio": next_monday.isoformat(),
        "periodo_fim": next_monday.isoformat(),
    }
    async with _client() as c:
        rec = await c.post(
            "/api/v1/admin/certificados",
            data=data,
            files={"documento": PDF},
            headers=_auth(aluno_admin["admin_id"]),
        )
        assert rec.status_code == 201, rec.text
        cert_id = rec.json()["id"]
        detail = await c.get(
            f"/api/v1/admin/certificados/{cert_id}", headers=_auth(aluno_admin["admin_id"])
        )
    refs = {d["turno_referencia"] for d in detail.json()["dias"]}
    assert f"{next_monday.isoformat()}_MANHA" in refs
